# مسیر: app/routers/template.py
"""
Template Download Router
~~~~~~~~~~~~~~~~~~~~~~~~
Dynamic Excel template generation for data import.

Route: /uploads/template/store
"""

from io import BytesIO

from fastapi import APIRouter
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

router = APIRouter(prefix="/uploads/template", tags=["template"])

# ─────────────────────────────────────────────────────
#  Column definitions — matched to Store model
# ─────────────────────────────────────────────────────

STORE_TEMPLATE_COLUMNS: list[dict] = [
    {
        "label": "نام فروشگاه",
        "column": "canonical_name",
        "required": True,
        "width": 30,
        "comment": "نام رسمی/ثبتی فروشگاه",
    },
    {
        "label": "آدرس",
        "column": "address",
        "required": True,
        "width": 50,
        "comment": "آدرس کامل پستی",
    },
    {
        "label": "تلفن",
        "column": "canonical_phone",
        "required": True,
        "width": 18,
        "comment": "شماره تماس (مثلاً 02112345678)",
    },
    {
        "label": "عرض جغرافیایی (Latitude)",
        "column": "latitude",
        "required": True,
        "width": 20,
        "comment": "مثلاً 35.6892",
    },
    {
        "label": "طول جغرافیایی (Longitude)",
        "column": "longitude",
        "required": True,
        "width": 20,
        "comment": "مثلاً 51.3890",
    },
    {
        "label": "دسته‌بندی کسب‌وکار",
        "column": "canonical_category",
        "required": False,
        "width": 22,
        "comment": "اختیاری — مثلاً سوپرمارکت",
    },
    {
        "label": "نوع فروشگاه",
        "column": "shop_type",
        "required": False,
        "width": 22,
        "comment": "اختیاری — سوپرمارکت / هایپرمارکت / فروشگاه زنجیره‌ای",
    },
    {
        "label": "کد پستی",
        "column": "postal_code",
        "required": False,
        "width": 15,
        "comment": "اختیاری — کد پستی ۱۰ رقمی",
    },
    {
        "label": "پلاک",
        "column": "plaque",
        "required": False,
        "width": 10,
        "comment": "اختیاری",
    },
    {
        "label": "واحد",
        "column": "unit",
        "required": False,
        "width": 10,
        "comment": "اختیاری",
    },
    {
        "label": "طبقه",
        "column": "floor",
        "required": False,
        "width": 10,
        "comment": "اختیاری",
    },
    {
        "label": "نام مدیر فروشگاه",
        "column": "manager_name",
        "required": False,
        "width": 25,
        "comment": "اختیاری — در CompanyStore ذخیره می‌شود",
    },
    {
        "label": "شماره موبایل",
        "column": "mobile",
        "required": False,
        "width": 18,
        "comment": "اختیاری — در CompanyStore ذخیره می‌شود",
    },
]

# ─────────────────────────────────────────────────────
#  Styles
# ─────────────────────────────────────────────────────

HEADER_FILL = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
HEADER_FONT = Font(name="B Nazanin", size=12, bold=True, color="FFFFFF")
REQUIRED_FILL = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
OPTIONAL_FILL = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
EXAMPLE_FONT = Font(name="B Nazanin", size=10, color="6B7280")
BODY_FONT = Font(name="B Nazanin", size=11)
THIN_BORDER = Border(
    left=Side(style="thin", color="D1D5DB"),
    right=Side(style="thin", color="D1D5DB"),
    top=Side(style="thin", color="D1D5DB"),
    bottom=Side(style="thin", color="D1D5DB"),
)
CENTER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
RIGHT_ALIGN = Alignment(horizontal="right", vertical="center", wrap_text=True)


@router.get("/store")
async def download_store_template():
    """
    Generate and serve a dynamic Excel template for store import.

    Columns match the Store model (app/models/store.py).
    Required columns are highlighted in blue.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "فروشگاه‌ها"
    ws.right_to_left = True

    # ── Row 1: Header labels ──────────────────────
    for col_idx, col_def in enumerate(STORE_TEMPLATE_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_def["label"])
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_BORDER

    # ── Row 2: DB column names (reference) ─────────
    for col_idx, col_def in enumerate(STORE_TEMPLATE_COLUMNS, start=1):
        cell = ws.cell(row=2, column=col_idx, value=col_def["column"])
        cell.font = Font(name="Consolas", size=9, color="9CA3AF", italic=True)
        cell.fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_BORDER

    # ── Row 3: Required/Optional marker ────────────
    for col_idx, col_def in enumerate(STORE_TEMPLATE_COLUMNS, start=1):
        marker = "✅ اجباری" if col_def["required"] else "⚪ اختیاری"
        cell = ws.cell(row=3, column=col_idx, value=marker)
        cell.font = Font(name="B Nazanin", size=9, bold=col_def["required"],
                         color="1E40AF" if col_def["required"] else "6B7280")
        cell.fill = REQUIRED_FILL if col_def["required"] else OPTIONAL_FILL
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_BORDER

    # ── Row 4: Example data ────────────────────────
    example_data = [
        "فروشگاه نمونه تهران",
        "تهران، خیابان انقلاب، نبش کوچه برهان",
        "02112345678",
        "35.6892",
        "51.3890",
        "سوپرمارکت",
        "سوپرمارکت",
        "1234567890",
        "۱۲",
        "۳",
        "همکف",
        "علی محمدی",
        "09121234567",
    ]
    for col_idx, value in enumerate(example_data, start=1):
        cell = ws.cell(row=4, column=col_idx, value=value)
        cell.font = EXAMPLE_FONT
        cell.alignment = RIGHT_ALIGN
        cell.border = THIN_BORDER

    # ── Row 5: Comment/Guide ──────────────────────
    comment_data = [
        "حتماً پر شود",
        "حتماً پر شود",
        "حتماً پر شود",
        "حتماً پر شود — عدد اعشاری",
        "حتماً پر شود — عدد اعشاری",
        "اختیاری",
        "اختیاری — سوپرمارکت/هایپرمارکت/فروشگاه زنجیره‌ای",
        "اختیاری — ۱۰ رقم",
        "اختیاری",
        "اختیاری",
        "اختیاری",
        "اختیاری",
        "اختیاری — ۱۱ رقم با ۰۹",
    ]
    for col_idx, value in enumerate(comment_data, start=1):
        cell = ws.cell(row=5, column=col_idx, value=value)
        cell.font = Font(name="B Nazanin", size=8, color="9CA3AF", italic=True)
        cell.alignment = RIGHT_ALIGN
        cell.border = THIN_BORDER

    # ── Column widths ─────────────────────────────
    for col_idx, col_def in enumerate(STORE_TEMPLATE_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = col_def["width"]

    # ── Row heights ───────────────────────────────
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 22
    ws.row_dimensions[4].height = 22
    ws.row_dimensions[5].height = 20

    # ── Freeze header rows ────────────────────────
    ws.freeze_panes = "A6"

    # ── Auto-filter ───────────────────────────────
    ws.auto_filter.ref = f"A1:{get_column_letter(len(STORE_TEMPLATE_COLUMNS))}5"

    # ── Stream to client ──────────────────────────
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": "attachment; filename=store_template.xlsx"
        }
    )
