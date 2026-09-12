"""Public access acceptance tests. Database targets are explicitly isolated."""
import io
import json
import os
import secrets
import zipfile
from datetime import timedelta
from types import SimpleNamespace
from uuid import uuid4
from unittest.mock import Mock

import pytest
from fastapi import FastAPI, HTTPException, Request
from app.services.audit.tests.demo_http_client import ASGIClient
from openpyxl import Workbook
from sqlalchemy import create_engine, text, select
from sqlalchemy.engine import make_url
from sqlalchemy.orm import Session
from app.models import Base, Account, Company, Project, File, ImportBatch
from app.models.demo_access import DemoChallenge, DemoSession
from app.core.database import get_db
from app.core.config import settings
from app.core.demo_security import DemoSecurityMiddleware, target
from app.core.trial_policy import policy
from app.repositories.demo_access_repository import DemoAccessRepository, digest, now
from app.services.demo_auth_service import DemoAuthService, LocalOtpProvider, normalize_mobile, code_hash
from app.services.trial_file_service import validate_trial_xlsx
from app.services.trial_service import TrialService
from app.services.demo_sample_service import public_sample_context


HEADERS = ['store_code', 'canonical_name', 'canonical_phone', 'address', 'latitude', 'longitude', 'manager_name', 'mobile']


def xlsx(columns=8, rows=1, headers=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Store Data'
    sheet.append(headers or (HEADERS + [f'extra{i}' for i in range(5)])[:columns])
    for i in range(rows):
        sheet.append(([str(i), f'Synthetic {i}', '02112345678', 'Synthetic Tehran street', 35.7, 51.4, '', ''] + [''] * 5)[:columns])
    payload = io.BytesIO()
    workbook.save(payload)
    return payload.getvalue()


@pytest.mark.parametrize('columns,rows', [(5,1),(8,1),(13,1),(8,1000)])
def test_xlsx_accepted_boundaries(columns, rows):
    assert validate_trial_xlsx('input.xlsx', xlsx(columns, rows)) == rows


@pytest.mark.parametrize('kind', ['14columns','1001rows','oversize','fake','corrupt','csv','multiple','macro','formula','external','zipbomb','encrypted','missing_name'])
def test_xlsx_rejected_boundaries(kind):
    payload, name = xlsx(), 'input.xlsx'
    if kind == '14columns':
        payload = xlsx(headers=HEADERS + ['a','b','c','d','e','f'])
    elif kind == '1001rows': payload = xlsx(rows=1001)
    elif kind == 'oversize': payload = b'PK\x03\x04' + b'x' * policy.file_bytes
    elif kind == 'fake': payload = b'not an excel file'
    elif kind == 'corrupt': payload = b'PK\x03\x04broken'
    elif kind == 'csv': name = 'input.csv'
    elif kind in {'multiple','formula','missing_name'}:
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(payload))
        if kind == 'multiple': book.create_sheet('Other')
        if kind == 'formula': book.active['A2'] = '=1+1'
        if kind == 'missing_name': book.active['B1'] = 'unknown'
        out = io.BytesIO(); book.save(out); payload = out.getvalue()
    elif kind == 'encrypted': payload = bytes.fromhex('D0CF11E0A1B11AE1') + b'encrypted'
    else:
        out = io.BytesIO(payload)
        with zipfile.ZipFile(out, 'a', compression=zipfile.ZIP_DEFLATED) as z:
            if kind == 'macro': z.writestr('xl/vbaProject.bin', b'macro')
            if kind == 'external': z.writestr('xl/externalLinks/link.xml', '<external/>')
            if kind == 'zipbomb': z.writestr('large.xml', b'x' * policy.expanded_bytes)
        payload = out.getvalue()
    with pytest.raises(HTTPException): validate_trial_xlsx(name, payload)


def test_xlsx_missing_coordinates_and_optional_fields_allowed():
    assert validate_trial_xlsx('input.xlsx', xlsx(columns=5, headers=['store_code','canonical_name','address','canonical_phone','extra'])) == 1


@pytest.mark.parametrize('mobile', ['09123456789','+989123456789','00989123456789','۹۱۲۳۴۵۶۷۸۹'])
def test_mobile_normalization(mobile):
    if mobile.startswith('۹'):
        with pytest.raises(HTTPException): normalize_mobile(mobile)
    else:
        assert normalize_mobile(mobile) == '09123456789'


@pytest.fixture(scope='module')
def engine():
    url = os.environ.get('DIP_DATASET_TEST_URL')
    if not url: pytest.skip('Explicit isolated test database required')
    parsed = make_url(url)
    assert parsed.host == 'dip-m21-db' and parsed.database == 'dip_m21_test'
    schema = 'demo_security_' + uuid4().hex[:12]
    root = create_engine(url)
    with root.begin() as c: c.execute(text(f'CREATE SCHEMA {schema}'))
    engine = create_engine(url, connect_args={'options': f'-csearch_path={schema},public'},
                           execution_options={'schema_translate_map': {None: schema}})
    with engine.begin() as c: Base.metadata.create_all(c)
    yield engine
    engine.dispose(); root.dispose()


@pytest.fixture
def db(engine):
    with Session(engine) as db:
        yield db
        db.rollback()


@pytest.fixture
def provider(monkeypatch):
    monkeypatch.setattr(settings, 'APP_ENV', 'development')
    monkeypatch.setenv('DEMO_OTP_PROVIDER', 'local')
    salt = secrets.token_hex(16)
    monkeypatch.setenv('DEMO_LOCAL_OTP_HASH', salt + ':' + code_hash('135790', salt))
    return LocalOtpProvider('testclient')


def identity(db):
    mobile = '09' + str(secrets.randbelow(10**9)).zfill(9)
    account = Account(mobile=mobile, mobile_verified_at=now())
    db.add(account); db.flush()
    company = Company(account_id=account.id, name='Synthetic tenant')
    db.add(company); db.flush()
    project = Project(company_id=company.id, trial_owner_id=account.id, name='Synthetic private project', trial_state='DRAFT')
    db.add(project); db.flush()
    token, csrf = secrets.token_hex(32), secrets.token_hex(32)
    session = DemoSession(account_id=account.id, token_hash=digest(token), csrf_hash=digest(csrf), expires_at=now()+timedelta(hours=1))
    db.add(session); db.commit()
    return account, project, token, csrf


@pytest.fixture
def client(engine, monkeypatch):
    monkeypatch.setattr(settings, 'DEMO_SAMPLE_PROJECT_ID', 0)
    from app.main import app as main_app
    app = FastAPI()
    app.router.routes.extend(main_app.router.routes)
    app.add_middleware(DemoSecurityMiddleware, engine=engine, raise_errors=True)
    def isolated_db(request: Request):
        if getattr(request.state, 'demo_db', None) is not None:
            yield request.state.demo_db
        else:
            with Session(engine) as db: yield db
    app.dependency_overrides[get_db] = isolated_db
    yield ASGIClient(app)


@pytest.mark.parametrize('path', ['/projects/new','/projects/','/projects/25','/uploads/file/1/download','/import/batches/1'])
def test_anonymous_pages_redirect(client, path):
    response = client.get(path, follow_redirects=False)
    assert response.status_code == 303
    assert response.headers['location'] == '/demo/login'


@pytest.mark.parametrize('path', ['/projects/new','/projects/api/','/uploads/project/25','/projects/25/matching/rerun'])
def test_anonymous_posts_denied(client, path):
    assert client.post(path).status_code == 401


def login_as(client, token, csrf):
    client.cookies.set('dip_session', token)
    client.cookies.set('dip_csrf', csrf)
    client.headers['X-CSRF-Token'] = csrf


def test_horizontal_ownership_and_legacy(client, db):
    a, pa, token, csrf = identity(db)
    b, pb, _, _ = identity(db)
    legacy = Project(company_id=pa.company_id, name='Private legacy sentinel')
    db.add(legacy); db.commit()
    login_as(client, token, csrf)
    listing = client.get('/projects/api/')
    assert listing.status_code == 200
    assert {p['id'] for p in listing.json()} == {pa.id}
    for path in [f'/projects/{pb.id}', f'/projects/{legacy.id}', f'/projects/{pb.id}/readiness', f'/projects/api/{pb.id}']:
        assert client.get(path).status_code == 404
    for path in [f'/uploads/project/{pb.id}', f'/projects/{pb.id}/matching/rerun', f'/projects/{pb.id}/readiness/refresh']:
        assert client.post(path).status_code == 404
    for step in ['', '/intake', '/validation', '/readiness']:
        assert client.get(f'/projects/{pa.id}{step}').status_code == 200


def test_csrf_expiry_and_invalid_session(client, db):
    a, p, token, csrf = identity(db)
    login_as(client, token, csrf)
    assert client.post('/projects/new', data={'name':'safe'}, headers={'X-CSRF-Token':'wrong'}).status_code == 403
    assert client.post('/projects/new', data={'name':'safe'}, headers={'Origin':'https://evil.example'}).status_code == 403
    session = db.scalar(select(DemoSession).where(DemoSession.token_hash == digest(token)))
    session.expires_at = now()-timedelta(seconds=1); db.commit()
    assert client.get('/projects/').status_code == 401
    client.cookies.set('dip_session', 'invalid')
    assert client.get('/projects/').status_code == 401


def test_otp_replay_resend_attempts_expiry(db, provider):
    service = DemoAuthService(db, provider)
    mobile = '09'+str(secrets.randbelow(10**9)).zfill(9)
    ref = service.issue(mobile, 'testclient')
    with pytest.raises(HTTPException) as resend: service.issue(mobile, 'testclient')
    assert resend.value.status_code == 429
    for _ in range(5):
        with pytest.raises(HTTPException): service.verify(ref, '000000', 'testclient')
    with pytest.raises(HTTPException): service.verify(ref, '135790', 'testclient')
    challenge = service.repo.challenge(ref)
    assert challenge.attempts == 5
    challenge.attempts = 0; challenge.expires_at = now()-timedelta(seconds=1); db.commit()
    with pytest.raises(HTTPException): service.verify(ref, '135790', 'testclient')
    challenge.expires_at = now()+timedelta(seconds=120); db.commit()
    token, _ = service.verify(ref, '135790', 'testclient')
    assert service.repo.session(token)
    with pytest.raises(HTTPException): service.verify(ref, '135790', 'testclient')


def test_local_provider_impossible_in_production(monkeypatch, provider):
    monkeypatch.setattr(settings, 'APP_ENV', 'production')
    with pytest.raises(HTTPException): LocalOtpProvider('127.0.0.1')
    monkeypatch.setattr(settings, 'APP_ENV', 'development')
    with pytest.raises(HTTPException): LocalOtpProvider('203.0.113.1')


def test_sample_allowlist_strips_sensitive_context():
    raw = {'readiness_summary': {'total_stores':1000, 'percentage':96.2, 'status':'ready_with_warning', 'mobile':'secret'}, 'project': {'id':1}, 'batch':{'id':28},
           'map_geojson_json': json.dumps({'features':[{'geometry':{'type':'Point','coordinates':[51.4,35.7]},
               'properties':{'mobile':'secret','title':'sensitive','store_id':99,'manager_name':'secret','readiness_status':'Ready'}}]}),
           'audit_run': {'raw':'secret'}}
    safe = public_sample_context(raw)
    assert safe['readiness_summary']['percentage'] == 96.2
    assert safe['readiness_summary']['status'] == 'ready_with_warning'
    serialized = json.dumps(safe)
    assert 'secret' not in serialized and 'sensitive' not in serialized
    assert 'store_id' not in serialized and '"batch"' not in serialized


def test_sample_public_without_internal_links_or_writes(client, monkeypatch):
    from app.routers import project
    monkeypatch.setattr(settings, 'DEMO_SAMPLE_PROJECT_ID', 1)
    render = Mock(return_value=SimpleNamespace(context={'batch': object(), 'readiness_summary': {'total_stores':1000,'ready_stores':962}, 'map_geojson_json':'{"features":[]}'}))
    monkeypatch.setattr(project, 'readiness_step', render)
    for _ in range(2):
        response = client.get('/demo/sample')
        assert response.status_code == 200
        for forbidden in ['href="/projects','href="/dashboard','method="post"','/matching/rerun','"store_id"']:
            assert forbidden not in response.text
    assert render.call_count == 2
    assert client.post('/projects/1/matching/rerun').status_code == 401


def test_idempotent_project_and_direct_quota_guard(client, db):
    a, p, token, csrf = identity(db)
    login_as(client, token, csrf)
    for _ in range(2):
        response = client.post('/projects/new', data={'name':'another'}, follow_redirects=False)
        assert response.status_code == 303
        assert response.headers['location'] == f'/projects/{p.id}'
    assert len(DemoAccessRepository(db).projects(a.id)) == 1
    p.trial_consumed_at = now(); db.commit()
    assert client.post('/projects/new', data={'name':'bypass'}).status_code == 403
    assert client.post('/projects/api/', data={'name':'bypass'}).status_code == 403
    assert client.delete(f'/projects/api/{p.id}').status_code == 403


@pytest.mark.parametrize('payload', [b'fake', b'PK\x03\x04broken', b'x'*(2*1024*1024+1)], ids=['fake','corrupt','oversize'])
def test_rejected_upload_has_no_batch_dataset_or_consumption(client, db, payload):
    a, p, token, csrf = identity(db)
    login_as(client, token, csrf)
    before = db.query(ImportBatch).count()
    response = client.post(f'/uploads/project/{p.id}', data={'category':'stores'}, files={'files':('input.xlsx',payload)})
    assert response.status_code in {400,413}
    db.expire_all()
    assert db.query(ImportBatch).count() == before
    assert db.get(Project,p.id).active_store_batch_id is None
    assert db.get(Project,p.id).trial_consumed_at is None


def test_unknown_internal_endpoints_and_raw_static_closed(client):
    for path in ['/admin', '/dashboard', '/request', '/import/upload', '/static/uploads/private.xlsx', '/openapi.json', '/docs']:
        assert client.get(path).status_code == 404


def test_owner_upload_matching_readiness_and_consumption(client, db, monkeypatch, tmp_path):
    import re
    from app.services.store_matching_service import StoreMatchingService
    monkeypatch.setattr(settings, 'UPLOAD_DIR', str(tmp_path))
    a, p, token, csrf = identity(db)
    login_as(client, token, csrf)
    payload = xlsx(rows=10)
    response = client.post(f'/uploads/project/{p.id}', data={'category':'stores'}, files={'files':('input.xlsx',payload)}, follow_redirects=False)
    assert response.status_code == 303
    db.expire_all(); p = db.get(Project,p.id)
    batch = p.active_store_batch_id
    assert db.get(ImportBatch,batch).imported_rows == 10
    response = client.post(f'/uploads/project/{p.id}', data={'category':'stores'}, files={'files':('input.xlsx',payload)}, follow_redirects=False)
    assert response.status_code == 303
    db.expire_all(); assert db.get(Project,p.id).active_store_batch_id == batch
    original = StoreMatchingService.run_for_batch
    calls = []
    def counted(self, *args, **kwargs):
        calls.append(1)
        return original(self, *args, **kwargs)
    monkeypatch.setattr(StoreMatchingService, 'run_for_batch', counted)
    for _ in range(2):
        result = client.post(f'/projects/{p.id}/matching/rerun', follow_redirects=False)
        assert result.status_code == 303
    assert len(calls) == 1
    db.expire_all(); assert db.get(Project,p.id).trial_consumed_at is None
    assert client.post(f'/demo/result/{p.id}/ack', data={'receipt':'forged'}).status_code == 403
    for _ in range(2):
        result = client.get(f'/projects/{p.id}/readiness')
        assert result.status_code == 200
        assert 'data-receipt=' in result.text
    assert len(calls) == 1
    db.expire_all(); assert db.get(Project,p.id).trial_consumed_at is None
    receipt = re.search(r'data-receipt="([^"]+)"', result.text)[1]
    assert client.post(f'/demo/result/{p.id}/ack', data={'receipt':receipt}).status_code == 200
    assert client.post(f'/demo/result/{p.id}/ack', data={'receipt':receipt}).status_code == 200
    db.expire_all(); assert db.get(Project,p.id).trial_consumed_at is not None
    assert client.post('/projects/new', data={'name':'second'}).status_code == 403


def test_processing_failure_rolls_back_and_does_not_consume(client, db, monkeypatch, tmp_path):
    from app.services.import_service import ImportService
    monkeypatch.setattr(settings, 'UPLOAD_DIR', str(tmp_path))
    a, p, token, csrf = identity(db)
    login_as(client, token, csrf)
    before = db.query(ImportBatch).count()
    original = ImportService.process_batch
    def fail_after_import(self, *args, **kwargs):
        original(self, *args, **kwargs)
        raise HTTPException(503, 'Synthetic internal failure')
    monkeypatch.setattr(ImportService, 'process_batch', fail_after_import)
    response = client.post(f'/uploads/project/{p.id}', data={'category':'stores'}, files={'files':('input.xlsx',xlsx())})
    assert response.status_code == 503
    db.expire_all()
    assert db.query(ImportBatch).count() == before
    assert db.get(Project,p.id).active_store_batch_id is None
    assert db.get(Project,p.id).trial_consumed_at is None


def test_account_operation_lock_blocks_concurrent_write(client, db, engine):
    a, p, token, csrf = identity(db)
    login_as(client, token, csrf)
    with Session(engine) as concurrent:
        assert DemoAccessRepository(concurrent).lock('trial:' + str(a.id))
        assert client.post('/projects/new', data={'name':'race'}).status_code == 409
        concurrent.rollback()


def test_otp_http_login_session_and_csrf(client, db, provider):
    import re
    page = client.get('/demo/login')
    assert page.status_code == 200
    csrf = client.cookies['dip_login_csrf']
    mobile = '09'+str(secrets.randbelow(10**9)).zfill(9)
    sent = client.post('/demo/otp/request', data={'mobile':mobile, 'csrf_token':csrf})
    assert sent.status_code == 200
    assert '135790' not in sent.text and mobile not in sent.text
    reference = re.search(r'name="reference" value="([^"]+)"', sent.text)[1]
    result = client.post('/demo/otp/verify', data={'reference':reference,'code':'135790','csrf_token':csrf}, follow_redirects=False)
    assert result.status_code == 303
    assert 'dip_session' in client.cookies
    assert client.get('/projects/new').status_code == 200
    result = client.post('/projects/new', data={'name':'My synthetic project','csrf_token':client.cookies['dip_csrf']}, follow_redirects=False)
    assert result.status_code == 303


def test_failed_matching_status_is_durable_and_retry_allowed(client, db, monkeypatch, tmp_path):
    from app.services.store_matching_service import StoreMatchingService
    monkeypatch.setattr(settings, 'UPLOAD_DIR', str(tmp_path))
    a, p, token, csrf = identity(db)
    login_as(client, token, csrf)
    response = client.post(f'/uploads/project/{p.id}', data={'category':'stores'}, files={'files':('input.xlsx',xlsx())}, follow_redirects=False)
    assert response.status_code == 303
    def failure(*args, **kwargs):
        raise HTTPException(503, 'Synthetic failure')
    monkeypatch.setattr(StoreMatchingService, 'run_for_batch', failure)
    assert client.post(f'/projects/{p.id}/matching/rerun').status_code == 503
    db.expire_all(); project = db.get(Project,p.id)
    assert project.trial_state == 'FAILED'
    assert project.trial_consumed_at is None
    assert project.active_store_batch_id


def test_files_and_batches_cannot_cross_owners(client, db, monkeypatch, tmp_path):
    monkeypatch.setattr(settings, 'UPLOAD_DIR', str(tmp_path))
    b, pb, tb, cb = identity(db)
    login_as(client,tb,cb)
    response = client.post(f'/uploads/project/{pb.id}', data={'category':'stores'}, files={'files':('input.xlsx',xlsx())}, follow_redirects=False)
    assert response.status_code == 303
    db.expire_all(); batch = db.get(ImportBatch, db.get(Project,pb.id).active_store_batch_id)
    file_id, batch_id = batch.file_id, batch.id
    a, pa, ta, ca = identity(db)
    login_as(client,ta,ca)
    for path in [f'/files/{file_id}/download', f'/uploads/file/{file_id}/download', f'/import/batches/{batch_id}', f'/import/api/batches/{batch_id}']:
        assert client.get(path).status_code == 404
    for path in [f'/files/{file_id}/delete', f'/uploads/file/{file_id}/delete']:
        assert client.post(path).status_code == 404


def test_rate_limit_and_redacted_logs(db, caplog):
    import logging
    from app.core.demo_logging import trial_request, install
    repo = DemoAccessRepository(db)
    key = 'synthetic-rate:' + uuid4().hex
    assert repo.rate(key, 2, 60)
    assert repo.rate(key, 2, 60)
    assert not repo.rate(key, 2, 60)
    install()
    token = trial_request.set(True)
    try:
        with caplog.at_level(logging.WARNING):
            logging.getLogger('app.services.import_service').warning('synthetic-private-mobile %s', 'synthetic-otp-secret')
    finally:
        trial_request.reset(token)
    assert 'synthetic-private-mobile' not in caplog.text
    assert 'synthetic-otp-secret' not in caplog.text
