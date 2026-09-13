import pytest
from sqlalchemy.orm import Session
from app.services.audit.tests.test_demo_security import engine, db, client, identity, login_as, xlsx
from app.repositories.demo_access_repository import DemoAccessRepository, now

HTML = {'Accept':'text/html'}


def test_quota_stays_html_and_api_stays_json(client, db, monkeypatch, tmp_path):
    from app.core.config import settings
    monkeypatch.setattr(settings, 'UPLOAD_DIR', str(tmp_path))
    account, project, token, csrf = identity(db)
    login_as(client, token, csrf)
    uploaded = client.post(f'/uploads/project/{project.id}', data={'category':'stores'}, files={'files':('input.xlsx',xlsx())}, follow_redirects=False)
    assert uploaded.status_code == 303
    project.trial_consumed_at = now(); db.commit()
    response = client.post(f'/projects/{project.id}/matching/rerun', headers=HTML)
    assert response.status_code == 403
    assert 'text/html' in response.headers['content-type']
    assert 'سهمیه ارزیابی رایگان این حساب مصرف شده است.' in response.text
    assert f'/projects/{project.id}/readiness' in response.text
    response = client.post('/projects/api/', headers=HTML, data={'name':'blocked'})
    assert response.status_code == 403 and 'application/json' in response.headers['content-type']


def test_concurrency_alert_preserves_409(client, db, engine):
    account, project, token, csrf = identity(db)
    login_as(client, token, csrf)
    with Session(engine) as locking:
        assert DemoAccessRepository(locking).lock('trial:'+str(account.id))
        response = client.post(f'/projects/{project.id}/readiness/refresh', headers=HTML)
        assert response.status_code == 409
        assert 'عملیات دیگری در حال اجراست.' in response.text
        assert 'data-web-error' in response.text


def test_forbidden_project_and_unknown_url_are_private_html(client, db):
    _, own, token, csrf = identity(db)
    _, other, _, _ = identity(db)
    login_as(client, token, csrf)
    for path in [f'/projects/{other.id}', '/۱/', '/not-a-page']:
        response = client.get(path, headers=HTML)
        assert response.status_code == 404
        assert 'امکان دسترسی به این صفحه وجود ندارد.' in response.text
        assert other.name not in response.text
        assert 'application/json' not in response.headers['content-type']
    assert client.get('/not-a-page').status_code == 404
    assert 'application/json' in client.get('/not-a-page').headers['content-type']


def test_session_csrf_and_rate_status_remain(client, db):
    _, project, token, csrf = identity(db)
    login_as(client, token, csrf)
    response = client.post(f'/projects/{project.id}/readiness/refresh', headers={**HTML,'X-CSRF-Token':'wrong'})
    assert response.status_code == 403 and 'data-web-error' in response.text
    client.cookies.set('dip_session', 'invalid')
    response = client.get('/projects/', headers=HTML)
    assert response.status_code == 401 and '/demo/login' in response.text


@pytest.mark.parametrize('status', [401,403,404,409,429,500])
def test_error_presentations_preserve_status(status):
    from starlette.requests import Request
    from app.core.web_errors import error_response
    for accept, expected in [('text/html','text/html'),('application/json','application/json')]:
        request = Request({'type':'http','method':'GET','path':'/missing','headers':[(b'accept',accept.encode())], 'scheme':'http','server':('testserver',80),'query_string':b''})
        response = error_response(request, status)
        assert response.status_code == status and expected in response.headers['content-type']


def test_invalid_upload_offers_authenticated_empty_existing_template(client, db):
    import io
    from openpyxl import load_workbook
    from app.models import Project, ImportBatch
    _, project, token, csrf = identity(db)
    login_as(client, token, csrf)
    before = db.query(ImportBatch).count()
    response = client.post(f'/uploads/project/{project.id}', headers=HTML, data={'category':'stores'}, files={'files':('invalid.xlsx', b'invalid')})
    assert response.status_code == 400
    assert 'برای آماده‌سازی فایل' in response.text
    assert 'href="/import/template/store"' in response.text
    downloaded = client.get('/import/template/store')
    assert downloaded.status_code == 200
    assert 'DIP_store_data_template.xlsx' in downloaded.headers['content-disposition']
    workbook = load_workbook(io.BytesIO(downloaded.content), read_only=True)
    assert len(workbook.worksheets) == 1 and workbook.active.max_column == 8
    assert all(v is None for row in workbook.active.iter_rows(min_row=2,values_only=True) for v in row)
    workbook.close()
    editable = load_workbook(io.BytesIO(downloaded.content))
    editable.active.cell(2,2,'Synthetic store')
    filled = io.BytesIO(); editable.save(filled); editable.close()
    from app.services.trial_file_service import validate_trial_xlsx
    assert validate_trial_xlsx('DIP_store_data_template.xlsx', filled.getvalue()) == 1
    db.expire_all()
    assert db.query(ImportBatch).count() == before
    assert db.get(Project, project.id).trial_consumed_at is None
    assert db.get(Project, project.id).active_store_batch_id is None
    client.cookies.clear(); client.headers.clear()
    assert client.get('/import/template/store', follow_redirects=False).status_code == 303


def test_same_mobile_different_code_preserves_account_project_and_consumption(client, db):
    from app.services.audit.tests.test_demo_codes import seed
    from app.models import Account, Project
    account, project, _, _ = identity(db)
    project.trial_consumed_at = now(); db.commit()
    consumed = project.trial_consumed_at
    account_count, project_count = db.query(Account).count(), db.query(Project).count()
    for code in [seed(db)[0], seed(db)[0]]:
        client.headers.clear()
        client.get('/demo/login')
        response = client.post('/demo/access', data={'mobile': account.mobile,'code':code,'csrf_token':client.cookies['dip_login_csrf']}, follow_redirects=False)
        assert response.status_code == 303 and response.headers['location'] == '/projects/'
        listing = client.get('/projects/api/').json()
        assert {item['id'] for item in listing} == {project.id}
        client.headers['X-CSRF-Token'] = client.cookies['dip_csrf']
        assert client.post('/projects/new',data={'name':'blocked'}).status_code == 403
        assert client.post(f'/uploads/project/{project.id}', data={'category':'stores'}, files={'files':('input.xlsx',xlsx())}).status_code == 403
    db.expire_all()
    assert db.query(Account).count() == account_count
    assert db.query(Project).count() == project_count
    assert db.get(Project, project.id).trial_consumed_at == consumed
