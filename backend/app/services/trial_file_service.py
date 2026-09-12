"""Validate bounded XLSX in memory before storage, ImportBatch or business writes."""
import io
import zipfile
from pathlib import PurePosixPath
from xml.etree import ElementTree as ET
from openpyxl import load_workbook
from fastapi import HTTPException
from app.core.trial_policy import policy
from app.services.import_service import ImportService
from app.schemas.import_schema import EntityTypeEnum


MESSAGE = "نسخه آزمایشی فقط فایل Excel با حداکثر ۱۰۰۰ ردیف، ۱۳ ستون و حجم ۲ مگابایت را می‌پذیرد."


def validate_trial_xlsx(filename, payload):
    try:
        if not filename.lower().endswith(".xlsx") or len(payload) > policy.file_bytes or not payload.startswith(b"PK\x03\x04"):
            raise ValueError()
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            entries = archive.infolist()
            names = [i.filename for i in entries]
            if len(entries) > policy.zip_entries or len(names) != len(set(names)):
                raise ValueError()
            if not {"[Content_Types].xml", "xl/workbook.xml", "_rels/.rels"} <= set(names):
                raise ValueError()
            if sum(i.file_size for i in entries) > policy.expanded_bytes:
                raise ValueError()
            for entry in entries:
                name = entry.filename.lower()
                if entry.is_dir():
                    continue
                if (entry.flag_bits & 1 or entry.file_size / max(entry.compress_size, 1) > policy.zip_ratio
                        or not name.endswith((".xml", ".rels"))
                        or ".." in PurePosixPath(name).parts or name.startswith("/")
                        or any(part in name for part in ("vbaproject", "macro", "embeddings/", "activex/", "externallinks/", "media/", "drawings/"))):
                    raise ValueError()
                content = archive.read(entry)
                if name.endswith((".xml", ".rels")):
                    normalized_xml = content.replace(b"\x00", b"").upper()
                    if b"<!DOCTYPE" in normalized_xml or b"<!ENTITY" in normalized_xml:
                        raise ValueError()
                    root = ET.fromstring(content)
                    for element in root.iter():
                        if element.tag.rsplit("}", 1)[-1] == "f" or element.attrib.get("TargetMode") == "External":
                            raise ValueError()
                        if "macroenabled" in element.attrib.get("ContentType", "").lower():
                            raise ValueError()
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=False, keep_links=False)
        try:
            if len(workbook.sheetnames) != 1:
                raise ValueError()
            sheet = workbook.active
            if sheet.max_row > policy.rows + 1 or not policy.min_columns <= sheet.max_column <= policy.max_columns:
                raise ValueError()
            # Do not trust a forged worksheet dimension to hide trailing rows.
            sheet.reset_dimensions()
            rows = []
            for row in sheet.iter_rows(values_only=True):
                if len(rows) > policy.rows or len(row) > policy.max_columns:
                    raise ValueError()
                rows.append(row)
            if len(rows) < 2 or len(rows) > policy.rows + 1:
                raise ValueError()
            headers = [str(cell or "").strip() for cell in rows[0]]
            if not policy.min_columns <= len(headers) <= policy.max_columns or len(set(headers)) != len(headers):
                raise ValueError()
            mapping = ImportService._suggest_mapping(headers, EntityTypeEnum.STORE)
            if len(headers) - len(set(mapping.values())) > policy.extra_columns:
                raise ValueError()
            # Existing import's identity mapping is authoritative; coordinates and
            # manager_name are optional. Never introduce an alternate alias table.
            if "canonical_name" not in mapping:
                raise HTTPException(400, "ستون نام فروشگاه در فایل قابل شناسایی نیست.")
            if not any(any(cell is not None for cell in row) for row in rows[1:]):
                raise ValueError()
            return len(rows) - 1
        finally:
            workbook.close()
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(400, MESSAGE) from None
